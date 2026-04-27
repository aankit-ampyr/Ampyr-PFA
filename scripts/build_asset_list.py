"""Build the canonical 78/120-slot asset list from F1 Project Info.

P2 deliverable per ADR 0001 (boolean flag), 0002 (JSON output), 0003 (all
120 slots with status enum).

Usage:
    python scripts/build_asset_list.py
    python scripts/build_asset_list.py --source path/to/F1.xlsm --out path/to/out.json

Reads:
    Project Info!I4:DX8  — across 120 horizontal slots:
        row 4  = Country (full)
        row 5  = Name
        row 6  = Technology
        row 7  = Include in Consolidation (boolean active flag — ADR 0001)
        row 8  = Country (Short form) — DE / UK / NL

Writes:
    data/asset_list.json — list of 120 dicts (one per slot), shape per ADR 0003.

Status enum (ADR 0003):
    - "active":      flag=True (in current scenario)
    - "placeholder": flag=False (per Anchal's mental model: ASE keeps a named pool of
                     inactive slots, replaces names when adding new real assets).
                     EMPIRICAL FINDING from F1: every one of 120 slots has a name —
                     there are no truly-empty slots. So "placeholder" here means
                     "named but inactive in current scenario."
    - "retired":     reserved for future use — requires F3-vs-F1 history comparison
                     (slots that flipped True→False between quarters). Bucket 2 #17
                     differ will populate this.
    - "unknown":     name set but flag is None (shouldn't occur in F1; defensive).
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).parent.parent
DEFAULT_SOURCE = ROOT / "data" / "converted" / \
    "Converted 2026-01-08 Project Parthenon - CIP_ASE_Internal_v18 - original file latest quarter.xlsm"
DEFAULT_OUT = ROOT / "data" / "asset_list.json"

# Slots are columns I (9) through DX (128) = 120 slots
SLOT_COL_START = 9
SLOT_COL_END = 128


def slot_status(name: str | None, is_active: bool | None) -> str:
    """Map (name, is_active) to slot_status enum per ADR 0003.

    Note: 'retired' is NOT inferred here from F1 alone. Distinguishing
    placeholder (inactive name reserved for future use) vs retired (was
    active, now turned off) requires comparing F1 to F3. The differ in
    Bucket 2 #17 will populate retired status. Until then all flag=False
    slots are tagged 'placeholder'.
    """
    if is_active is True:
        return "active"
    if is_active is False:
        return "placeholder"
    return "unknown"


def build(source: Path, out: Path) -> dict:
    if not source.exists():
        raise FileNotFoundError(source)

    wb = load_workbook(source, data_only=True, read_only=True)
    ws = wb["Project Info"]

    # Read rows 4, 5, 6, 7, 8 across cols I:DX
    def read_row(r: int) -> list:
        return list(next(ws.iter_rows(min_row=r, max_row=r,
                                       min_col=SLOT_COL_START,
                                       max_col=SLOT_COL_END,
                                       values_only=True)))

    countries_full = read_row(4)
    names = read_row(5)
    technologies = read_row(6)
    active_flags = read_row(7)
    countries_short = read_row(8)
    wb.close()

    slots = []
    for i in range(SLOT_COL_END - SLOT_COL_START + 1):
        slot_index = i + 1  # 1-based to match Excel slot numbering
        name = names[i]
        flag = active_flags[i]
        status = slot_status(name, flag)
        slots.append({
            "slot_index": slot_index,
            "name": name if (name is not None and str(name).strip() != "") else None,
            "country_short": countries_short[i],
            "country_full": countries_full[i],
            "technology": technologies[i],
            "is_active": flag if flag in (True, False) else None,
            "slot_status": status,
        })

    # Counts for the metadata block + sanity output
    counts = {"active": 0, "placeholder": 0, "retired": 0, "unknown": 0}
    for s in slots:
        counts[s["slot_status"]] = counts.get(s["slot_status"], 0) + 1

    payload = {
        "version_source": source.name,
        "extracted_at": datetime.now(timezone.utc).isoformat(),
        "extracted_by": "scripts/build_asset_list.py",
        "slot_count": len(slots),
        "status_counts": counts,
        "schema_version": 1,
        "decisions_applied": ["ADR-0001", "ADR-0002", "ADR-0003"],
        "slots": slots,
    }

    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")
    return payload


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()

    payload = build(args.source, args.out)
    print(f"Built {args.out.relative_to(ROOT)}")
    print(f"  source: {args.source.name}")
    print(f"  slots: {payload['slot_count']}")
    print(f"  status counts: {payload['status_counts']}")

    # Show 3 actives + 3 retired + 3 placeholders for quick eyeball
    print()
    print("Sample (first 3 of each status):")
    for status in ("active", "retired", "placeholder", "unknown"):
        matches = [s for s in payload["slots"] if s["slot_status"] == status]
        if matches:
            print(f"  --- {status} ({len(matches)}) ---")
            for s in matches[:3]:
                print(f"    slot {s['slot_index']:3d}: "
                      f"name={s['name']!r:35} "
                      f"country={s['country_short']!r:6} "
                      f"tech={s['technology']!r:20} "
                      f"flag={s['is_active']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
