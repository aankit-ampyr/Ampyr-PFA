"""06 - GTC sheets deep dive (F2 only):
- For each of 21 GTC sheets, list: cells, formulas, functions used, sheets referenced
- Build dependency graph: which GTC sheet pulls from which ASE sheet / GTC sheet
"""
import json
import re
from collections import Counter, defaultdict
from openpyxl import load_workbook
from pathlib import Path

DATA = Path("c:/repos/Ampyr-PFA/data/converted")
OUT = Path("c:/repos/Ampyr-PFA/.claude/analysis_2026_04")

PATH = DATA / "Converted FY27 Business Planv8 copy - file with GTC sheets added.xlsm"

FN_RE = re.compile(r"\b([A-Z][A-Z0-9._]*)\s*\(")
SHEET_REF_RE = re.compile(r"'([^']+)'!|\b([A-Za-z_][\w]*)!")

# GTC sheets added in F2 (not in F1)
GTC_SHEETS = [
    "Instruction tab", "Summary sheet", "PnL projection - aggregate",
    "FY26A-29F PL Breakdown", "FY26B PL Mapping", "BS projection - aggregate",
    "FY26A-29F BS Breakdown", "FY26B BS Mapping", "CF Capital Strategy",
    "CF Capital Strategy Breakdown", "Q Rep (USD)", "Valuation (2)",
    "Cash Breakeven", "QRep(EUR)", "FX", "Net income bridge",
    "Valuation", "Log", "Graphs", "Tables", "Asset workings",
]

wb = load_workbook(PATH, read_only=True, data_only=False, keep_links=False)

gtc_report = {}
for sheet in GTC_SHEETS:
    if sheet not in wb.sheetnames:
        print(f"  !! missing: {sheet}")
        continue
    ws = wb[sheet]
    fn = Counter()
    sheet_refs = Counter()
    n_cells = n_f = 0
    for row in ws.iter_rows():
        for c in row:
            v = getattr(c, "value", None)
            if v is None:
                continue
            n_cells += 1
            if isinstance(v, str) and v.startswith("="):
                n_f += 1
                body = v[1:]
                for f in FN_RE.findall(body):
                    fn[f] += 1
                for m in SHEET_REF_RE.findall(body):
                    s = m[0] if m[0] else m[1]
                    if s:
                        sheet_refs[s] += 1
    gtc_report[sheet] = {
        "cells": n_cells, "formulas": n_f,
        "fn_top": fn.most_common(10),
        "refs_top": sheet_refs.most_common(10),
    }
    print(f"  {sheet:40} cells={n_cells:>8} f={n_f:>8} fn:{fn.most_common(3)} refs:{sheet_refs.most_common(3)}")

wb.close()
(OUT / "06_gtc.json").write_text(json.dumps(gtc_report, indent=2, default=str))

# Build dependency graph
print("\nDependency matrix (GTC -> source sheet, counts of formula refs):")
sources = set()
for v in gtc_report.values():
    for s, _ in v["refs_top"]:
        sources.add(s)
sources = sorted(sources)
print(f"  sources: {sources[:15]}...")
print(f"\n  GTC sheet -> top ASE/GTC source sheets with counts:")
for gtc in GTC_SHEETS:
    if gtc in gtc_report:
        refs = dict(gtc_report[gtc]["refs_top"])
        top = sorted(refs.items(), key=lambda kv: -kv[1])[:5]
        print(f"    {gtc:40} -> {top}")
