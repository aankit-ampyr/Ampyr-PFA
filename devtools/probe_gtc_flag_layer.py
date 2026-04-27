"""Deep dive: where (if anywhere) does the GTC override flag live?

Anchal's mental model: ASE provides one flag, GTC modifies a second flag
"internally for assessments". Question is whether (i) GTC edits the same
row 7 cells in F2, (ii) there's a separate column/row, or (iii) it's
elsewhere (Sensis or a GTC-only sheet).

Steps:
1. Diff row 5 (asset names) and row 7 (active flags) between F1 and F2
2. List F2-only named ranges that might be flag-related
3. Scan F2's 21 added sheets for boolean / active / include columns
4. Look at Sensis for any per-asset toggles

Run: python devtools/probe_gtc_flag_layer.py
"""
from __future__ import annotations
from openpyxl import load_workbook
from pathlib import Path
import json

DATA = Path(__file__).parent.parent / "data" / "converted"
F1_PATH = DATA / "Converted 2026-01-08 Project Parthenon - CIP_ASE_Internal_v18 - original file latest quarter.xlsm"
F2_PATH = DATA / "Converted FY27 Business Planv8 copy - file with GTC sheets added.xlsm"
ANALYSIS = Path(__file__).parent.parent / ".claude" / "analysis_2026_04"


def step1_row7_diff():
    print("=" * 70)
    print("STEP 1: Diff row 5 (names) and row 7 (active flags) F1 vs F2")
    print("=" * 70)
    wb1 = load_workbook(F1_PATH, data_only=True, read_only=True)
    wb2 = load_workbook(F2_PATH, data_only=True, read_only=True)
    ws1 = wb1["Project Info"]
    ws2 = wb2["Project Info"]

    # ProjectActiveFlag is row 7 cols I:DX (cols 9..128 = 120 slots)
    r5_f1 = list(next(ws1.iter_rows(min_row=5, max_row=5, min_col=9, max_col=128, values_only=True)))
    r5_f2 = list(next(ws2.iter_rows(min_row=5, max_row=5, min_col=9, max_col=128, values_only=True)))
    r7_f1 = list(next(ws1.iter_rows(min_row=7, max_row=7, min_col=9, max_col=128, values_only=True)))
    r7_f2 = list(next(ws2.iter_rows(min_row=7, max_row=7, min_col=9, max_col=128, values_only=True)))

    diffs = []
    for slot, (n1, n2, f1, f2) in enumerate(zip(r5_f1, r5_f2, r7_f1, r7_f2), start=1):
        if n1 != n2 or f1 != f2:
            diffs.append((slot, n1, n2, f1, f2))

    n_active_f1 = sum(1 for f in r7_f1 if f is True)
    n_active_f2 = sum(1 for f in r7_f2 if f is True)
    print(f"  F1 active count: {n_active_f1}  |  F2 active count: {n_active_f2}")
    print(f"  Slots where row5 (name) or row7 (flag) differ: {len(diffs)}")
    for slot, n1, n2, f1, f2 in diffs[:10]:
        print(f"    slot {slot}: name F1={n1!r} F2={n2!r} | flag F1={f1!r} F2={f2!r}")
    if len(diffs) > 10:
        print(f"    ... ({len(diffs) - 10} more)")
    wb1.close()
    wb2.close()
    return diffs


def step2_named_ranges():
    print()
    print("=" * 70)
    print("STEP 2: F2-only named ranges with flag-like names")
    print("=" * 70)
    nr = json.loads((ANALYSIS / "02_named_ranges.json").read_text())
    f1_names = {n["name"] for n in nr["F1"].get("wb_named_ranges", [])}
    f2_names = {n["name"]: n.get("value", "") for n in nr["F2"].get("wb_named_ranges", [])}
    f2_only = {nm: ref for nm, ref in f2_names.items() if nm not in f1_names}

    keywords = ["active", "consolid", "flag", "include", "exclud", "select", "override", "live"]
    matches = {nm: ref for nm, ref in f2_only.items()
               if any(k in nm.lower() for k in keywords)}
    print(f"  F2-only names: {len(f2_only)}")
    print(f"  Flag-like (filtered): {len(matches)}")
    for nm in sorted(matches.keys()):
        print(f"    {nm:40} → {matches[nm][:80]}")

    # Also check sheet-scoped names in F2 (entries may be dicts or strings)
    sheet_named = nr["F2"].get("sheet_scoped_names", [])

    def _name_of(item) -> str:
        return item.get("name", "") if isinstance(item, dict) else str(item)

    sheet_matches = [s for s in sheet_named
                     if any(k in _name_of(s).lower() for k in keywords)]
    print(f"  Sheet-scoped flag-like names in F2: {len(sheet_matches)}")
    for s in sheet_matches[:15]:
        if isinstance(s, dict):
            print(f"    {s.get('name', ''):40} on {s.get('sheet', '?')[:25]} → {s.get('value', '')[:60]}")
        else:
            print(f"    {s}")


def step3_scan_added_sheets():
    print()
    print("=" * 70)
    print("STEP 3: Scan F2's added sheets for active/include columns")
    print("=" * 70)
    inv = json.loads((ANALYSIS / "01_inventory.json").read_text())
    f1_sheets = {s["name"] for s in inv["F1_ASE_Latest"].get("sheets", [])}
    f2_sheets = {s["name"] for s in inv["F2_GTC_Enhanced"].get("sheets", [])}
    added = sorted(f2_sheets - f1_sheets)
    print(f"  Sheets added in F2: {len(added)}")

    wb2 = load_workbook(F2_PATH, data_only=True, read_only=True)
    keywords = ["active", "include", "exclud", "consolidat", "override", "selected", "in_scope"]

    for name in added:
        ws = wb2[name]
        # Scan first 5 rows for header keywords (cap on cols to avoid huge sheets)
        max_col = min(ws.max_column, 200)
        hits = []
        for row in ws.iter_rows(min_row=1, max_row=5, max_col=max_col, values_only=True):
            for col_idx, val in enumerate(row, start=1):
                if val is None:
                    continue
                s = str(val).lower()
                if any(k in s for k in keywords):
                    hits.append((col_idx, str(val)[:50]))
        if hits:
            print(f"  '{name}': {len(hits)} flag-like header(s)")
            for col, val in hits[:5]:
                print(f"    col {col}: {val!r}")
    wb2.close()


def step4_sensis():
    print()
    print("=" * 70)
    print("STEP 4: Sensis — any per-asset toggles?")
    print("=" * 70)
    wb1 = load_workbook(F1_PATH, data_only=True, read_only=True)
    if "Sensis" not in wb1.sheetnames:
        print("  No Sensis sheet in F1")
        wb1.close()
        return
    ws = wb1["Sensis"]
    print(f"  Sensis dim: {ws.max_row}r x {ws.max_column}c")

    # Look for rows with many bool cells (indicating per-asset toggles)
    bool_rich_rows = []
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 100),
                                              max_col=min(ws.max_column, 150),
                                              values_only=True), start=1):
        n_bool = sum(1 for v in row if isinstance(v, bool))
        if n_bool >= 10:  # 10+ bools in a row = suspicious
            bool_rich_rows.append((r_idx, n_bool, row[1] if len(row) > 1 else None))
    print(f"  Rows with 10+ booleans (potential per-asset toggle): {len(bool_rich_rows)}")
    for r, n, label in bool_rich_rows[:10]:
        print(f"    row {r}: {n} bool cells, label={label!r}")

    wb1.close()


def main():
    step1_row7_diff()
    step2_named_ranges()
    step3_scan_added_sheets()
    step4_sensis()


if __name__ == "__main__":
    main()
