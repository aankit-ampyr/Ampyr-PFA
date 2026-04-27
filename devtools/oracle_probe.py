"""Tiny validation-oracle probe: do the workbooks contain cached values?"""
from openpyxl import load_workbook
import time, sys

F1 = 'data/converted/Converted 2026-01-08 Project Parthenon - CIP_ASE_Internal_v18 - original file latest quarter.xlsm'

def probe(path, sheet, rows, cols, label):
    t0 = time.time()
    wb = load_workbook(path, data_only=True, read_only=True)
    print(f'  load {label}: {time.time()-t0:.1f}s', flush=True)
    ws = wb[sheet]
    n_pop = n_none = 0
    samples = []
    # iter_rows is the read_only fast path
    for r_idx, row in enumerate(ws.iter_rows(min_row=rows[0], max_row=rows[1],
                                              min_col=cols[0], max_col=cols[1],
                                              values_only=True), start=rows[0]):
        for c_idx, v in enumerate(row, start=cols[0]):
            if v is None:
                n_none += 1
            else:
                n_pop += 1
                if len(samples) < 3:
                    samples.append((r_idx, c_idx, type(v).__name__, str(v)[:30]))
    total = n_pop + n_none
    pct = 100*n_pop/total if total else 0
    print(f'  {sheet}: {n_pop}/{total} populated ({pct:.0f}%)', flush=True)
    for s in samples:
        print(f'    sample {s}', flush=True)
    wb.close()
    return pct

print('=== F1 PLW row 250-260, col AB-AZ (deep in time-axis numeric region) ===', flush=True)
probe(F1, 'Project Level Workings', (250, 260), (28, 52), 'F1')
print(flush=True)
print('=== F1 Quarterly Output row 10-110, col B-AZ ===', flush=True)
probe(F1, 'Quarterly Output', (10, 110), (2, 52), 'F1-cached')
